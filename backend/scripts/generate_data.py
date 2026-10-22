"""Generate the Summit Supply Co. sample database.

The output is fully deterministic (fixed seed), so evaluation results are
reproducible. The data is synthetic but shaped like a real business:

  * the customer base grows over time; orders peak in November / December and
    business customers rarely order on weekends;
  * some customers never order, some churn, a few buy far more than the rest;
  * an order's status follows its age at the snapshot date (AS_OF): recent
    orders are still Pending / Processing / Shipped, older ones are Delivered,
    with a small share Cancelled or Returned;
  * corporate customers pay by bank transfer on credit terms, so recent
    invoices are still unpaid and a handful are overdue;
  * a few newly launched products have never sold, discontinued products stop
    selling, and order lines carry the historical price, not today's list price.

Usage (from backend/):
    python scripts/generate_data.py                  # writes data/company.db
    python scripts/generate_data.py --xlsx           # also writes data/company_sample.xlsx
"""

from __future__ import annotations

import argparse
import math
import random
import re
import sqlite3
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

from faker import Faker

from agent.knowledge.us_geography import STATES, STATES_BY_CODE

SEED = 42
AS_OF = date(2026, 10, 20)  # snapshot date: "today" from the data's point of view
COMPANY_START = date(2014, 3, 3)
CUSTOMERS_START = date(2021, 6, 1)
ORDERS_START = date(2022, 1, 1)
N_CUSTOMERS = 3000
EMAIL_DOMAIN = "summitsupply.com"

BACKEND_DIR = Path(__file__).resolve().parents[1]
DATA_DIR = BACKEND_DIR / "data"
SCHEMA_PATH = DATA_DIR / "schema.sql"

rng = random.Random(SEED)
Faker.seed(SEED)
fake = Faker("en_US")
fake.seed_instance(SEED)


# ─── small helpers ───────────────────────────────────────────────────────────


def rand_date(start: date, end: date) -> date:
    return start + timedelta(days=rng.randint(0, max(0, (end - start).days)))


def pick(options: dict):
    """Weighted choice from {value: weight}."""
    return rng.choices(list(options), weights=list(options.values()), k=1)[0]


def poisson(lam: float) -> int:
    if lam <= 0:
        return 0
    if lam > 30:
        return max(0, round(rng.gauss(lam, math.sqrt(lam))))
    limit, k, p = math.exp(-lam), 0, 1.0
    while True:
        p *= rng.random()
        if p <= limit:
            return k
        k += 1


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", text.lower())


def iso(d: date | None) -> str | None:
    return d.isoformat() if d else None


STATE_CODES = [s.code for s in STATES]
STATE_WEIGHTS = [s.population_m for s in STATES]


def random_state() -> str:
    return rng.choices(STATE_CODES, weights=STATE_WEIGHTS, k=1)[0]


def random_city(state_code: str) -> str:
    cities = STATES_BY_CODE[state_code].cities
    # Larger (first-listed) cities are more likely.
    return rng.choices(cities, weights=[1 / (i + 1) ** 0.6 for i in range(len(cities))], k=1)[0]


def phone() -> str:
    return f"({rng.randint(201, 989)}) {rng.randint(200, 999)}-{rng.randint(0, 9999):04d}"


_used_emails: set[str] = set()


def unique_email(local: str, domain: str) -> str:
    local = re.sub(r"[^a-z0-9.]", "", local.lower()).strip(".")
    candidate, n = f"{local}@{domain}", 1
    while candidate in _used_emails:
        n += 1
        candidate = f"{local}{n}@{domain}"
    _used_emails.add(candidate)
    return candidate


# ─── employees ───────────────────────────────────────────────────────────────

OFFICES = {
    "San Francisco": "CA",
    "New York": "NY",
    "Austin": "TX",
    "Dallas": "TX",
    "Chicago": "IL",
    "Seattle": "WA",
}


@dataclass
class Employee:
    first: str
    last: str
    title: str
    department: str
    hire: date
    salary: float
    city: str
    state: str
    kind: str  # exec | head | manager | ic
    employment_type: str = "Full-Time"
    status: str = "Active"
    termination: date | None = None
    manager: Employee | None = None
    id: int = 0

    def active_on(self, d: date) -> bool:
        return self.hire <= d and (self.termination is None or self.termination > d)


@dataclass(frozen=True)
class Dept:
    name: str
    office: str
    owner: str  # title of the executive the department head reports to
    head: tuple[str, int, int]
    manager: tuple[str, int, int]
    n_managers: int
    n_ics: int
    roles: tuple[tuple[str, int, int, float], ...]  # (title, salary_lo, salary_hi, weight)
    part_time_share: float = 0.0
    contract_share: float = 0.0
    spread_offices: bool = False


EXECUTIVES = (
    ("Chief Executive Officer", 340_000, 340_000),
    ("Chief Financial Officer", 265_000, 290_000),
    ("Chief Operating Officer", 255_000, 285_000),
    ("Chief Technology Officer", 275_000, 300_000),
    ("Chief Revenue Officer", 260_000, 290_000),
)

DEPARTMENTS = (
    Dept("Sales", "New York", "Chief Revenue Officer", ("VP of Sales", 210_000, 235_000),
         ("Sales Manager", 125_000, 155_000), 6, 80,
         (("Sales Development Representative", 55_000, 72_000, 3), ("Sales Representative", 68_000, 95_000, 5),
          ("Senior Sales Representative", 90_000, 125_000, 3), ("Account Executive", 95_000, 140_000, 4)),
         spread_offices=True),
    Dept("Marketing", "San Francisco", "Chief Revenue Officer", ("VP of Marketing", 195_000, 220_000),
         ("Marketing Manager", 115_000, 140_000), 3, 22,
         (("Marketing Specialist", 65_000, 90_000, 3), ("Content Strategist", 70_000, 95_000, 2),
          ("Digital Marketing Analyst", 72_000, 100_000, 2), ("Brand Designer", 75_000, 105_000, 1)),
         contract_share=0.1),
    Dept("Engineering", "Seattle", "Chief Technology Officer", ("VP of Engineering", 235_000, 260_000),
         ("Engineering Manager", 185_000, 220_000), 6, 64,
         (("Software Engineer", 115_000, 160_000, 5), ("Senior Software Engineer", 150_000, 205_000, 4),
          ("Staff Engineer", 195_000, 245_000, 1), ("Data Engineer", 125_000, 175_000, 2),
          ("QA Engineer", 90_000, 130_000, 2), ("DevOps Engineer", 125_000, 170_000, 2)),
         contract_share=0.1),
    Dept("Product", "San Francisco", "Chief Technology Officer", ("Head of Product", 215_000, 240_000),
         ("Group Product Manager", 180_000, 205_000), 2, 14,
         (("Product Manager", 130_000, 175_000, 3), ("Senior Product Manager", 160_000, 205_000, 2),
          ("Product Designer", 115_000, 160_000, 2), ("UX Researcher", 105_000, 145_000, 1)),
         contract_share=0.1),
    Dept("Customer Support", "Austin", "Chief Operating Officer", ("Director of Customer Support", 135_000, 155_000),
         ("Support Manager", 80_000, 98_000), 4, 50,
         (("Support Specialist", 45_000, 60_000, 6), ("Senior Support Specialist", 55_000, 72_000, 3),
          ("Support Team Lead", 65_000, 82_000, 1)),
         part_time_share=0.15),
    Dept("Operations", "Dallas", "Chief Operating Officer", ("Director of Operations", 145_000, 165_000),
         ("Operations Manager", 90_000, 115_000), 3, 38,
         (("Warehouse Associate", 38_000, 52_000, 5), ("Logistics Coordinator", 52_000, 70_000, 3),
          ("Inventory Analyst", 60_000, 82_000, 2), ("Procurement Specialist", 62_000, 85_000, 2)),
         part_time_share=0.08),
    Dept("Finance", "New York", "Chief Financial Officer", ("Controller", 175_000, 195_000),
         ("Finance Manager", 120_000, 145_000), 2, 15,
         (("Financial Analyst", 80_000, 115_000, 3), ("Senior Accountant", 78_000, 105_000, 2),
          ("Accountant", 62_000, 82_000, 2), ("Payroll Specialist", 58_000, 76_000, 1))),
    Dept("Human Resources", "Chicago", "Chief Operating Officer", ("Director of People", 150_000, 170_000),
         ("HR Manager", 95_000, 118_000), 1, 10,
         (("HR Generalist", 60_000, 82_000, 2), ("Recruiter", 65_000, 95_000, 3),
          ("HR Business Partner", 85_000, 120_000, 1))),
)

# Company growth: more hires in recent years (2026 is a partial year).
IC_HIRE_YEAR_WEIGHTS = {2016: 1, 2017: 2, 2018: 3, 2019: 4, 2020: 3, 2021: 6, 2022: 8,
                        2023: 8, 2024: 10, 2025: 11, 2026: 8}


def salary_in(lo: int, hi: int, factor: float = 1.0) -> float:
    return round(rng.uniform(lo, hi) * factor / 500) * 500


def new_person(title: str, dept: str, hire: date, salary: float, city: str, kind: str) -> Employee:
    return Employee(fake.first_name(), fake.last_name(), title, dept, hire, salary, city, OFFICES.get(city, ""), kind)


def generate_employees() -> list[Employee]:
    staff: list[Employee] = []
    execs: dict[str, Employee] = {}

    for i, (title, lo, hi) in enumerate(EXECUTIVES):
        hire = COMPANY_START if i == 0 else rand_date(date(2014, 4, 1), date(2017, 12, 31))
        e = new_person(title, "Executive", hire, salary_in(lo, hi), "San Francisco", "exec")
        execs[title] = e
        staff.append(e)
    ceo = execs["Chief Executive Officer"]
    for e in execs.values():
        if e is not ceo:
            e.manager = ceo

    sales_offices = {"New York": 4, "San Francisco": 3, "Austin": 2, "Chicago": 2, "Seattle": 1}

    for dept in DEPARTMENTS:
        head = new_person(dept.head[0], dept.name, rand_date(date(2015, 1, 1), date(2020, 6, 30)),
                          salary_in(dept.head[1], dept.head[2]), dept.office, "head")
        head.manager = execs[dept.owner]
        staff.append(head)

        managers = []
        for _ in range(dept.n_managers):
            office = pick(sales_offices) if dept.spread_offices else dept.office
            m = new_person(dept.manager[0], dept.name, rand_date(date(2016, 1, 1), date(2023, 6, 30)),
                           salary_in(dept.manager[1], dept.manager[2]), office, "manager")
            m.manager = head
            managers.append(m)
        staff.extend(managers)

        role_weights = {r[0]: r[3] for r in dept.roles}
        role_bands = {r[0]: (r[1], r[2]) for r in dept.roles}
        for _ in range(dept.n_ics):
            year = pick(IC_HIRE_YEAR_WEIGHTS)
            hire = rand_date(date(year, 1, 1), min(date(year, 12, 31), AS_OF - timedelta(days=3)))
            title = pick(role_weights)
            emp_type = pick({"Full-Time": 1 - dept.part_time_share - dept.contract_share,
                             "Part-Time": dept.part_time_share, "Contract": dept.contract_share})
            factor = 0.55 if emp_type == "Part-Time" else 1.0
            office = pick(sales_offices) if dept.spread_offices else dept.office
            e = new_person(title, dept.name, hire, salary_in(*role_bands[title], factor), office, "ic")
            e.employment_type = emp_type
            if rng.random() < 0.12:  # remote employee
                e.state = random_state()
                e.city = random_city(e.state)
            senior_enough = [m for m in managers if m.hire <= hire]
            e.manager = rng.choice(senior_enough) if senior_enough else head
            staff.append(e)

    for e in staff:
        if e.kind != "ic":
            continue
        tenure_days = (AS_OF - e.hire).days
        if tenure_days > 150 and rng.random() < 0.15:
            e.status = "Terminated"
            e.termination = e.hire + timedelta(days=rng.randint(120, tenure_days - 1))
        elif rng.random() < 0.03:
            e.status = "On Leave"

    staff.sort(key=lambda e: (e.hire, e.last, e.first))
    for i, e in enumerate(staff, start=1):
        e.id = i
        e.email = unique_email(f"{e.first}.{e.last}", EMAIL_DOMAIN)
    return staff


# ─── products ────────────────────────────────────────────────────────────────

CATALOG = {
    # category: (sku prefix, price range, brands, base product names)
    "Electronics": ("ELEC", (25, 420), ("Voltix", "Auralis", "Nexon"),
                    ("Wireless Earbuds", "Bluetooth Speaker", "Noise-Cancelling Headphones", "Smartwatch",
                     "4K Streaming Stick", "Portable Charger", "Smart Home Hub", "Action Camera", "E-Reader",
                     "Fitness Tracker")),
    "Computers & Accessories": ("COMP", (18, 1900), ("Corelink", "Pixelon", "Datamax"),
                                ("14-inch Laptop", "15-inch Laptop", "27-inch Monitor", "Mechanical Keyboard",
                                 "Wireless Mouse", "USB-C Hub", "1TB External SSD", "1080p Webcam",
                                 "Laptop Stand", "Docking Station", "11-inch Tablet")),
    "Office Supplies": ("OFFC", (4, 260), ("Paperline", "Officio"),
                        ("Printer Paper Case", "Ballpoint Pens 24-Pack", "Sticky Notes 12-Pack", "Heavy-Duty Stapler",
                         "Desk Organizer", "Whiteboard Markers 8-Pack", "File Folders 100-Pack", "Label Maker",
                         "Cross-Cut Shredder", "Toner Cartridge")),
    "Furniture": ("FURN", (45, 1300), ("ErgoForm", "Oakridge"),
                  ("Ergonomic Office Chair", "Standing Desk", "Bookshelf", "Filing Cabinet", "Conference Table",
                   "LED Desk Lamp", "Dual Monitor Arm", "Guest Chair")),
    "Home & Kitchen": ("HOME", (20, 460), ("KitchenPro", "Homely"),
                       ("Coffee Maker", "Air Fryer", "High-Speed Blender", "Cookware Set", "Chef Knife Set",
                        "Electric Kettle", "Cordless Vacuum", "Water Filter Pitcher", "Toaster Oven")),
    "Sports & Outdoors": ("SPRT", (10, 360), ("TrailPeak", "FitCore"),
                          ("Yoga Mat", "Adjustable Dumbbells", "Camping Tent", "Hiking Backpack",
                           "Insulated Water Bottle", "Resistance Bands Set", "Bike Helmet", "Sleeping Bag")),
    "Books": ("BOOK", (10, 60), ("Lumen Press", "Northbound Books"),
              ("Data Science Handbook", "Leadership Essentials", "Python Programming Guide", "Personal Finance 101",
               "Mystery Novel Collection", "Weeknight Meals Cookbook", "History of Innovation",
               "Mindfulness Workbook")),
    "Health & Personal Care": ("HLTH", (8, 210), ("Vitalis", "CareWell"),
                               ("Electric Toothbrush", "Digital Thermometer", "First Aid Kit", "Ionic Hair Dryer",
                                "Massage Gun", "Vitamin D Supplements", "Blood Pressure Monitor")),
}

# Which categories each segment tends to buy.
CATEGORY_AFFINITY = {
    "Consumer": {"Electronics": 1.4, "Computers & Accessories": 1.0, "Office Supplies": 0.5, "Furniture": 0.4,
                 "Home & Kitchen": 1.4, "Sports & Outdoors": 1.1, "Books": 1.0, "Health & Personal Care": 1.1},
    "Small Business": {"Electronics": 0.7, "Computers & Accessories": 1.5, "Office Supplies": 1.8, "Furniture": 1.0,
                       "Home & Kitchen": 0.4, "Sports & Outdoors": 0.2, "Books": 0.4, "Health & Personal Care": 0.3},
    "Corporate": {"Electronics": 0.6, "Computers & Accessories": 1.8, "Office Supplies": 1.6, "Furniture": 1.3,
                  "Home & Kitchen": 0.3, "Sports & Outdoors": 0.1, "Books": 0.3, "Health & Personal Care": 0.2},
}


@dataclass
class Product:
    id: int
    name: str
    sku: str
    category: str
    brand: str
    price: float
    cost: float
    launch: date
    discontinued_on: date | None
    popularity: float
    stock: int = 0
    reorder: int = 0

    def sellable_on(self, d: date) -> bool:
        return self.launch <= d and (self.discontinued_on is None or d < self.discontinued_on)


def list_price(lo: float, hi: float) -> float:
    # Log-uniform inside the category's range, ending in .99
    value = math.exp(rng.uniform(math.log(lo), math.log(hi)))
    return max(lo, math.floor(value)) + 0.99


def generate_products() -> list[Product]:
    products: list[Product] = []
    counters: dict[str, int] = {}
    for category, (prefix, (lo, hi), brands, bases) in CATALOG.items():
        for base in bases:
            n_brands = 2 if rng.random() < 0.55 else 1
            for brand in rng.sample(brands, n_brands):
                counters[prefix] = counters.get(prefix, 0) + 1
                price = list_price(lo, hi)
                cost = round(price * rng.uniform(0.42, 0.72), 2)
                launch = (rand_date(date(2019, 1, 1), date(2021, 12, 31)) if rng.random() < 0.7
                          else rand_date(date(2022, 1, 1), date(2026, 8, 31)))
                products.append(Product(0, f"{brand} {base}", f"{prefix}-{counters[prefix]:04d}", category,
                                        brand, price, cost, launch, None, rng.paretovariate(1.6)))

    # ~9% of older products were discontinued at some point in 2024-2025.
    for p in products:
        if p.launch < date(2023, 1, 1) and rng.random() < 0.09:
            p.discontinued_on = rand_date(date(2024, 1, 1), date(2025, 12, 31))

    # Four brand-new launches that have not sold yet ("products never ordered").
    for p in rng.sample([p for p in products if p.discontinued_on is None], 4):
        p.launch = rand_date(date(2026, 10, 1), date(2026, 10, 18))
        p.popularity = 0.0

    products.sort(key=lambda p: (p.category, p.name))
    for i, p in enumerate(products, start=1):
        p.id = i
        p.reorder = rng.choice((10, 15, 20, 25, 30, 40, 50, 75, 100))
        if p.discontinued_on:
            p.stock = rng.randint(0, 12)
        elif rng.random() < 0.12:  # running low
            p.stock = rng.randint(0, p.reorder)
        else:
            p.stock = rng.randint(p.reorder + 1, p.reorder * 12)
    return products


# ─── customers ───────────────────────────────────────────────────────────────


@dataclass
class Customer:
    first: str
    last: str
    company: str | None
    email: str
    phone: str | None
    city: str
    state: str
    postal: str
    segment: str
    signup: date
    account_manager: Employee | None
    propensity: float
    churn: date | None
    never_buys: bool
    tier: str = "Bronze"
    id: int = 0


def signup_date() -> date:
    # Accept-reject so later dates (a growing customer base) are more likely.
    span = (AS_OF - CUSTOMERS_START).days
    while True:
        t = rng.random()
        if rng.random() < (1 + 1.5 * t) / 2.5:
            return CUSTOMERS_START + timedelta(days=int(t * span))


def postal_code(state: str) -> str:
    try:
        return fake.zipcode_in_state(state)
    except Exception:  # Faker does not cover every code (e.g. DC in some versions)
        return fake.zipcode()


def generate_customers(employees: list[Employee]) -> list[Customer]:
    account_managers = [e for e in employees if e.department == "Sales" and e.status != "Terminated"
                        and e.title in ("Account Executive", "Senior Sales Representative", "Sales Representative")]
    customers: list[Customer] = []
    for _ in range(N_CUSTOMERS):
        segment = pick({"Consumer": 0.66, "Small Business": 0.22, "Corporate": 0.12})
        first, last = fake.first_name(), fake.last_name()
        state = random_state()
        if segment == "Consumer":
            company = None
            email = unique_email(f"{first}.{last}{rng.randint(1, 99)}", fake.free_email_domain())
            manager = None
        else:
            company = fake.company()
            email = unique_email(f"{first}.{last}", f"{slug(company)[:20]}.com")
            manager = rng.choice(account_managers)
        signup = signup_date()
        churn = None
        if rng.random() < 0.22 and (AS_OF - signup).days > 120:
            churn = signup + timedelta(days=rng.randint(60, (AS_OF - signup).days))
        customers.append(Customer(
            first, last, company, email, phone() if rng.random() > 0.09 else None,
            random_city(state), state, postal_code(state), segment, signup, manager,
            rng.lognormvariate(0, 0.8) / 1.377, churn, rng.random() < 0.08,
        ))
    customers.sort(key=lambda c: (c.signup, c.last, c.first))
    for i, c in enumerate(customers, start=1):
        c.id = i
    return customers


# ─── orders ──────────────────────────────────────────────────────────────────

MONTH_SEASONALITY = {1: 0.8, 2: 0.8, 3: 0.95, 4: 0.95, 5: 1.0, 6: 1.0, 7: 0.95, 8: 1.0, 9: 1.0,
                     10: 1.05, 11: 1.45, 12: 1.6}
MONTHLY_ORDER_RATE = {"Consumer": 0.17, "Small Business": 0.33, "Corporate": 0.5}


@dataclass
class OrderLine:
    product: Product
    quantity: int
    unit_price: float
    discount: float

    @property
    def net(self) -> float:
        return self.quantity * self.unit_price * (1 - self.discount)


@dataclass
class Order:
    customer: Customer
    order_date: date
    lines: list[OrderLine] = field(default_factory=list)
    employee: Employee | None = None
    status: str = "Pending"
    channel: str = "Online"
    payment_method: str = "Credit Card"
    payment_status: str = "Paid"
    shipping_cost: float = 0.0
    shipped: date | None = None
    delivered: date | None = None
    ship_city: str = ""
    ship_state: str = ""
    id: int = 0


def order_day_weight(d: date, segment: str) -> float:
    weekend = d.weekday() >= 5
    dow = (1.1 if weekend else 1.0) if segment == "Consumer" else (0.25 if weekend else 1.0)
    return MONTH_SEASONALITY[d.month] * dow


def sample_order_date(start: date, end: date, segment: str) -> date:
    while True:
        d = rand_date(start, end)
        if rng.random() < order_day_weight(d, segment) / 1.76:
            return d


def order_dates_for(c: Customer) -> list[date]:
    if c.never_buys:
        return []
    start = max(c.signup, ORDERS_START)
    end = min(c.churn or AS_OF, AS_OF)
    if end <= start:
        return []
    months = (end - start).days / 30.44
    n = poisson(MONTHLY_ORDER_RATE[c.segment] * c.propensity * months)
    dates = []
    if n and c.signup >= ORDERS_START and rng.random() < 0.6:
        first = c.signup + timedelta(days=min(int(rng.expovariate(1 / 10)), 60))
        if first <= end:
            dates.append(first)
            n -= 1
    dates.extend(sample_order_date(start, end, c.segment) for _ in range(max(0, n)))
    return dates


def quantity_for(segment: str, category: str) -> int:
    if segment == "Consumer":
        q = pick({1: 75, 2: 17, 3: 6, 4: 2})
        return q * (rng.randint(1, 3) if category in ("Office Supplies", "Books") else 1)
    if category == "Office Supplies":
        hi = 60 if segment == "Corporate" else 20
    elif category in ("Computers & Accessories", "Furniture"):
        hi = 25 if segment == "Corporate" else 8
    else:
        hi = 12 if segment == "Corporate" else 5
    return max(1, int(rng.triangular(1, hi, 1)))


def discount_for(segment: str, d: date) -> float:
    if segment == "Consumer":
        if d.month in (11, 12) and rng.random() < 0.35:
            return rng.choice((0.10, 0.15, 0.20, 0.25))
        return rng.choice((0.05, 0.10, 0.15)) if rng.random() < 0.15 else 0.0
    if rng.random() < 0.45:
        return pick({0.05: 3, 0.10: 4, 0.15: 3, 0.20: 2, 0.25: 1, 0.30: 0.5})
    return 0.0


def historical_price(p: Product, d: date) -> float:
    years_back = (AS_OF - d).days / 365.25
    return max(0.5, round(p.price * (1 - 0.03 * years_back) * rng.uniform(0.97, 1.03), 2))


def build_lines(o: Order, products: list[Product]) -> None:
    seg = o.customer.segment
    eligible = [p for p in products if p.sellable_on(o.order_date) and p.popularity > 0]
    weights = [p.popularity * CATEGORY_AFFINITY[seg][p.category] for p in eligible]
    max_lines = 8 if seg == "Consumer" else 12
    n_lines = min(max_lines, 1 + int(rng.expovariate(0.8)), len(eligible))
    chosen: list[Product] = []
    while len(chosen) < n_lines:
        p = rng.choices(eligible, weights=weights, k=1)[0]
        if p not in chosen:
            chosen.append(p)
    for p in chosen:
        o.lines.append(OrderLine(p, quantity_for(seg, p.category), historical_price(p, o.order_date),
                                 discount_for(seg, o.order_date)))


def handler_for(o: Order, employees: list[Employee]) -> Employee | None:
    """The employee who booked the order (NULL for self-service online consumer orders)."""
    c = o.customer
    if c.segment != "Consumer":
        if c.account_manager and c.account_manager.active_on(o.order_date):
            return c.account_manager
    elif o.channel == "Online":
        return None
    reps = [e for e in employees if e.department == "Sales" and e.kind == "ic" and e.active_on(o.order_date)]
    return rng.choice(reps) if reps else None


def finalize_fulfilment(o: Order) -> None:
    c = o.customer
    business = c.segment != "Consumer"
    subtotal = sum(line.net for line in o.lines)

    # Destination: usually the customer's own address, sometimes another state (gift / branch office).
    if rng.random() < 0.9:
        o.ship_city, o.ship_state = c.city, c.state
    else:
        o.ship_state = random_state()
        o.ship_city = random_city(o.ship_state)

    if o.channel == "In-Store":
        o.shipping_cost = 0.0
    elif business:
        o.shipping_cost = 0.0 if subtotal >= 500 else 19.99
    else:
        o.shipping_cost = 0.0 if subtotal >= 75 else rng.choice((5.99, 7.99, 9.99))

    if c.segment == "Consumer":
        o.payment_method = pick({"Credit Card": 0.64, "PayPal": 0.2, "Digital Wallet": 0.1, "Gift Card": 0.06})
    else:
        bank = 0.65 if c.segment == "Corporate" else 0.45
        o.payment_method = pick({"Bank Transfer": bank, "Credit Card": 0.93 - bank, "PayPal": 0.07})

    cancelled = rng.random() < (0.03 if business else 0.05)
    returned = rng.random() < (0.02 if business else 0.05)
    has_furniture = any(line.product.category == "Furniture" for line in o.lines)

    if o.channel == "In-Store":
        ship_lag, deliver_lag = 0, 0
    else:
        ship_lag = rng.randint(1, 3) + (rng.randint(2, 5) if has_furniture else 0)
        deliver_lag = rng.randint(2, 6)
    shipped = o.order_date + timedelta(days=ship_lag)
    delivered = shipped + timedelta(days=deliver_lag)

    if cancelled:
        o.status = "Cancelled"
        o.payment_status = "Voided" if o.payment_method == "Bank Transfer" or rng.random() < 0.55 else "Refunded"
        return

    if o.order_date >= AS_OF:
        o.status = "Pending"
    elif shipped > AS_OF:
        o.status = "Processing" if o.order_date < AS_OF - timedelta(days=1) or rng.random() < 0.5 else "Pending"
    elif delivered > AS_OF:
        o.status, o.shipped = "Shipped", shipped
    else:
        o.status, o.shipped, o.delivered = "Delivered", shipped, delivered
        if returned and delivered + timedelta(days=10) <= AS_OF:
            o.status = "Returned"

    if o.status == "Returned":
        o.payment_status = "Refunded"
    elif o.payment_method == "Bank Transfer":
        # Net-terms invoices: paid 5-45 days after ordering; ~2% end up overdue.
        paid_on = o.order_date + timedelta(days=rng.randint(5, 45))
        o.payment_status = "Paid" if paid_on <= AS_OF and rng.random() > 0.02 else "Pending"
    elif o.status == "Pending" and rng.random() < 0.5:
        o.payment_status = "Pending"
    else:
        o.payment_status = "Paid"


def generate_orders(customers: list[Customer], employees: list[Employee], products: list[Product]) -> list[Order]:
    orders: list[Order] = []
    for c in customers:
        for d in order_dates_for(c):
            o = Order(c, d)
            o.channel = pick({"Online": 0.72, "Phone": 0.08, "In-Store": 0.12, "Partner": 0.08}
                             if c.segment == "Consumer"
                             else {"Online": 0.35, "Phone": 0.30, "In-Store": 0.10, "Partner": 0.25})
            orders.append(o)
    orders.sort(key=lambda o: (o.order_date, o.customer.id))
    for i, o in enumerate(orders, start=1):
        o.id = i
        build_lines(o, products)
        o.employee = handler_for(o, employees)
        finalize_fulfilment(o)
    return orders


def assign_loyalty_tiers(customers: list[Customer], orders: list[Order]) -> None:
    spend: dict[int, float] = {}
    for o in orders:
        if o.status == "Delivered":
            spend[o.customer.id] = spend.get(o.customer.id, 0.0) + sum(line.net for line in o.lines)
    for c in customers:
        s = spend.get(c.id, 0.0)
        c.tier = "Platinum" if s >= 25_000 else "Gold" if s >= 5_000 else "Silver" if s >= 1_000 else "Bronze"


# ─── persistence ─────────────────────────────────────────────────────────────


def write_database(path: Path, employees, customers, products, orders) -> None:
    if path.exists():
        path.unlink()
    con = sqlite3.connect(path)
    try:
        con.executescript(SCHEMA_PATH.read_text(encoding="utf-8"))
        # IDs follow hire date, so a manager hired after one of their reports
        # (e.g. a replacement manager) is inserted later: check FKs at COMMIT.
        con.execute("BEGIN")
        con.execute("PRAGMA defer_foreign_keys = ON")
        con.executemany(
            "INSERT INTO Employees VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            [(e.id, e.first, e.last, e.email, e.title, e.department, e.manager.id if e.manager else None,
              iso(e.hire), iso(e.termination), e.salary, e.employment_type, e.status, e.city, e.state)
             for e in employees],
        )
        con.executemany(
            "INSERT INTO Customers VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
            [(c.id, c.first, c.last, c.company, c.email, c.phone, c.city, c.state, c.postal, c.segment, c.tier,
              iso(c.signup), c.account_manager.id if c.account_manager else None) for c in customers],
        )
        con.executemany(
            "INSERT INTO Products VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            [(p.id, p.name, p.sku, p.category, p.brand, p.price, p.cost, p.stock, p.reorder,
              1 if p.discontinued_on else 0, iso(p.launch)) for p in products],
        )
        con.executemany(
            "INSERT INTO Orders VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
            [(o.id, o.customer.id, o.employee.id if o.employee else None, iso(o.order_date), o.status, o.channel,
              o.payment_method, o.payment_status, o.shipping_cost, iso(o.shipped), iso(o.delivered),
              o.ship_city, o.ship_state) for o in orders],
        )
        item_rows, item_id = [], 0
        for o in orders:
            for line in o.lines:
                item_id += 1
                item_rows.append((item_id, o.id, line.product.id, line.quantity, line.unit_price, line.discount))
        con.executemany("INSERT INTO OrderItems VALUES (?,?,?,?,?,?)", item_rows)
        con.commit()

        problems = con.execute("PRAGMA foreign_key_check").fetchall()
        if problems:
            raise RuntimeError(f"foreign key violations: {problems[:5]}")
        if con.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise RuntimeError("integrity check failed")
        con.execute("ANALYZE")  # planner statistics, used for query cost estimation
        con.commit()
        con.execute("VACUUM")
    finally:
        con.close()


def write_workbook(db_path: Path, xlsx_path: Path) -> None:
    """Export every table to one sheet each, so the data can be browsed in Excel."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    con = sqlite3.connect(db_path)
    wb = Workbook(write_only=True)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    tables = [r[0] for r in con.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY rowid")]
    for table in tables:
        cur = con.execute(f'SELECT * FROM "{table}"')
        headers = [d[0] for d in cur.description]
        ws = wb.create_sheet(table)
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(32, len(h) + 4))
        cells = []
        for h in headers:
            cell = WriteOnlyCell(ws, value=h)
            cell.font, cell.fill = header_font, header_fill
            cells.append(cell)
        ws.append(cells)
        for row in cur:
            ws.append(list(row))
    con.close()
    wb.save(xlsx_path)


def print_summary(db_path: Path) -> None:
    con = sqlite3.connect(db_path)
    q = lambda sql: con.execute(sql).fetchall()  # noqa: E731
    print(f"\nDatabase written: {db_path}  (snapshot as of {AS_OF})")
    for (t,) in q("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY rowid"):
        print(f"  {t:<11} {q(f'SELECT COUNT(*) FROM {t}')[0][0]:>7,} rows")
    print("\nNet revenue by year (excl. cancelled/returned):")
    for year, revenue, n in q("""
        SELECT strftime('%Y', o.OrderDate), ROUND(SUM(oi.Quantity * oi.UnitPrice * (1 - oi.Discount)), 0),
               COUNT(DISTINCT o.OrderID)
        FROM Orders o JOIN OrderItems oi ON oi.OrderID = o.OrderID
        WHERE o.Status NOT IN ('Cancelled', 'Returned') GROUP BY 1 ORDER BY 1"""):
        print(f"  {year}: ${revenue:>13,.0f}  across {n:,} orders")
    print("\nOrder status mix:", dict(q("SELECT Status, COUNT(*) FROM Orders GROUP BY 1 ORDER BY 2 DESC")))
    print("Loyalty tiers:  ", dict(q("SELECT LoyaltyTier, COUNT(*) FROM Customers GROUP BY 1 ORDER BY 2 DESC")))
    print("Top customer states:", q("SELECT State, COUNT(*) FROM Customers GROUP BY 1 ORDER BY 2 DESC LIMIT 5"))
    print("Customers with no orders:",
          q("SELECT COUNT(*) FROM Customers WHERE CustomerID NOT IN (SELECT CustomerID FROM Orders)")[0][0])
    print("Products never ordered:  ",
          q("SELECT COUNT(*) FROM Products WHERE ProductID NOT IN (SELECT ProductID FROM OrderItems)")[0][0])
    print("Employees hired since 2024-01-01:",
          q("SELECT COUNT(*) FROM Employees WHERE HireDate >= '2024-01-01'")[0][0])
    con.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", type=Path, default=DATA_DIR / "company.db")
    parser.add_argument("--xlsx", nargs="?", const=DATA_DIR / "company_sample.xlsx", type=Path, default=None,
                        help="also export the tables to an Excel workbook")
    args = parser.parse_args()

    employees = generate_employees()
    products = generate_products()
    customers = generate_customers(employees)
    orders = generate_orders(customers, employees, products)
    assign_loyalty_tiers(customers, orders)

    write_database(args.out, employees, customers, products, orders)
    print_summary(args.out)
    if args.xlsx:
        write_workbook(args.out, args.xlsx)
        print(f"\nWorkbook written: {args.xlsx}")


if __name__ == "__main__":
    main()
